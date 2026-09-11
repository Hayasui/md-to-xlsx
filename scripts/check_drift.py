# -*- coding: utf-8 -*-
"""判断交付表在生成之后有没有被人手动改过（模板）。

思路：加载项目的生成脚本，把它的 `META["filename"]` 换成临时路径重建一份，
再与磁盘上的正式文件逐格比对；差异格就是别人的手动改动。
项目脚本若把输出路径写死，在这里改 `META` 即可，不动原脚本。

用法：把 PROJ / BUILD 改成项目实际值（写在源码里，别走命令行参数——
中文路径当 argv 传给 python 在本机偶发编码错乱），然后 `python check_drift.py`，
报告写到同目录 `_drift.txt`（UTF-8）。
"""
import os, sys, importlib.util, zipfile
from openpyxl import load_workbook

# ↓↓↓ 按项目改 ↓↓↓
PROJ = r"C:\path\to\project"
BUILD = os.path.join(PROJ, ".workbuddy", "build", "build_project_xlsx.py")
TMP = os.path.join(PROJ, ".workbuddy", "build", "_rebuilt.xlsx")
OUT = os.path.join(PROJ, ".workbuddy", "build", "_drift.txt")
# ↑↑↑ 按项目改 ↑↑↑

# 加载项目脚本但不执行它的 __main__，好拿到里面的 META 与建表函数
spec = importlib.util.spec_from_file_location("proj_build", BUILD)
b = importlib.util.module_from_spec(spec)
spec.loader.exec_module(b)

target = b.META["filename"]
saved = b.META["filename"]
b.META["filename"] = TMP
sys.argv = [BUILD]
try:
    b.main()                       # 项目脚本若已挂上技能的 main()，这里会连带跑一次校验
finally:
    b.META["filename"] = saved

a = load_workbook(target)
c = load_workbook(TMP)
out = ["现有 sheets: %s" % a.sheetnames, "重建 sheets: %s" % c.sheetnames]
total = 0
for name in a.sheetnames:
    if name not in c.sheetnames:
        out.append("[!] 现有表 %s 在重建版里不存在" % name)
        continue
    sa, sc = a[name], c[name]
    if (sa.max_row, sa.max_column) != (sc.max_row, sc.max_column):
        out.append("[!] %s 尺寸不同：现有 %dx%d / 重建 %dx%d"
                   % (name, sa.max_row, sa.max_column, sc.max_row, sc.max_column))
    diffs = 0
    for r in range(1, max(sa.max_row, sc.max_row) + 1):
        for col in range(1, max(sa.max_column, sc.max_column) + 1):
            va = sa.cell(r, col).value or ""
            vc = sc.cell(r, col).value or ""
            if va != vc:
                diffs += 1
                if diffs <= 30:
                    out.append("  [%s] R%dC%d\n    现有: %r\n    重建: %r"
                               % (name, r, col, va, vc))
    out.append("%s：差异 %d 处" % (name, diffs))
    total += diffs
out.append("合计差异：%d（0 表示生成后没被人动过）" % total)

# 顺带看文件被什么工具保存过
with zipfile.ZipFile(target) as z:
    try:
        app = z.read("docProps/app.xml").decode("utf-8", "ignore")
        out.append("docProps: %s" % app[:300])
    except KeyError:
        out.append("docProps/app.xml 缺失")

with open(OUT, "w", encoding="utf-8") as fh:
    fh.write("\n".join(out))
print("报告写到:", OUT)
